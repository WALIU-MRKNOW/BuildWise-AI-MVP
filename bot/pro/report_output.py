"""In-memory BuildWise Pro report output renderers.

The report remains a draft engineering-assistance document in every format. No output is
persisted by this module.
"""

from __future__ import annotations

from html import escape
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from bot.pro.models import PROFESSIONAL_REVIEW_NOTICE, ProfessionalReport


def _report_sections(report: ProfessionalReport) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    if report.template_structure_used and report.template_sections:
        sections.extend((section.heading, list(section.content)) for section in report.template_sections)
    else:
        sections.extend(
            [
                ("Project / document information", list(report.project_document_information)),
                ("Executive summary", [report.summary]),
                ("Observations", list(report.observations)),
                ("Drawing / sketch findings", list(report.drawing_sketch_findings)),
                ("Site-note findings", list(report.site_note_findings)),
                ("Technical considerations", list(report.relevant_technical_considerations)),
                ("Standard / code reference", list(report.referenced_code_standard)),
                ("Recommendations", list(report.recommendations)),
            ]
        )
    sections.extend(
        [
            ("Limitations / assumptions", list(report.limitations_assumptions)),
            ("Uncertainties / missing information", list(report.uncertainties_missing_information)),
            ("Professional review / sign-off", [report.review_signoff]),
        ]
    )
    return sections


def render_report_html(report: ProfessionalReport, project_information: str) -> bytes:
    """Render the existing self-contained HTML fallback."""

    def list_html(items: list[str]) -> str:
        if not items:
            return "<p>Insufficient information provided.</p>"
        return "<ul>" + "".join(f"<li>{escape(item)}</li>" for item in items) + "</ul>"

    body = "".join(
        f"<section><h2>{escape(title)}</h2>{list_html(items)}</section>"
        for title, items in _report_sections(report)
    )
    html = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>BuildWise Pro Draft Report</title>
<style>body{{font-family:Arial,sans-serif;max-width:900px;margin:40px auto;padding:0 24px;line-height:1.55;color:#171717}}h1,h2{{line-height:1.2}}.notice{{border:3px solid #9b1c1c;padding:16px;margin:20px 0;font-weight:700}}section{{margin:28px 0}}.meta{{color:#555}}</style></head>
<body><h1>BuildWise Pro — Draft Professional Report</h1>
<p class="meta">Project input: {escape(project_information[:2000])}</p>
<div class="notice">{escape(PROFESSIONAL_REVIEW_NOTICE)}</div>
{body}
<div class="notice">{escape(PROFESSIONAL_REVIEW_NOTICE)}</div>
</body></html>"""
    return html.encode("utf-8")


def render_report_pdf(report: ProfessionalReport, project_information: str) -> bytes:
    """Render a portable PDF draft entirely in memory."""
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="BuildWise Pro Draft Professional Report",
        author="BuildWise Pro",
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.alignment = TA_CENTER
    notice_style = ParagraphStyle(
        "BuildWiseNotice",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#8A1111"),
    )
    body_style = styles["BodyText"]
    body_style.leading = 14
    heading_style = styles["Heading2"]
    heading_style.spaceBefore = 10
    heading_style.spaceAfter = 6

    story: list[object] = [
        Paragraph("BuildWise Pro — Draft Professional Report", title_style),
        Spacer(1, 5 * mm),
        Paragraph(f"<b>Project input:</b> {escape(project_information[:2000])}", body_style),
        Spacer(1, 4 * mm),
    ]

    def notice_box() -> Table:
        table = Table([[Paragraph(escape(PROFESSIONAL_REVIEW_NOTICE), notice_style)]], colWidths=[170 * mm])
        table.setStyle(
            TableStyle(
                [
                    ("BOX", (0, 0), (-1, -1), 1.3, colors.HexColor("#8A1111")),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF5F5")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        return table

    story.extend([notice_box(), Spacer(1, 5 * mm)])
    for heading, items in _report_sections(report):
        story.append(Paragraph(escape(heading), heading_style))
        if not items:
            story.append(Paragraph("Insufficient information provided.", body_style))
        elif heading in {"Executive summary", "Professional review / sign-off"}:
            for item in items:
                story.append(Paragraph(escape(item), body_style))
                story.append(Spacer(1, 2 * mm))
        else:
            bullets = [
                ListItem(Paragraph(escape(item), body_style), leftIndent=10)
                for item in items
            ]
            story.append(ListFlowable(bullets, bulletType="bullet", leftIndent=18))
        story.append(Spacer(1, 2 * mm))
    story.extend([Spacer(1, 4 * mm), notice_box()])
    document.build(story)
    return buffer.getvalue()


def render_report_docx(report: ProfessionalReport, project_information: str) -> bytes:
    """Render an editable DOCX draft entirely in memory."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = document.add_heading("BuildWise Pro — Draft Professional Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.add_run("Project input: ").bold = True
    meta.add_run(project_information[:2000])

    def add_notice() -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(PROFESSIONAL_REVIEW_NOTICE)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = None
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(10)

    add_notice()
    for heading, items in _report_sections(report):
        document.add_heading(heading, level=1)
        if not items:
            document.add_paragraph("Insufficient information provided.")
        elif heading in {"Executive summary", "Professional review / sign-off"}:
            for item in items:
                document.add_paragraph(item)
        else:
            for item in items:
                document.add_paragraph(item, style="List Bullet")
    add_notice()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
