from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock

from docx import Document

from bot.config import Settings
from bot.pro.ai_service import ProFile, ProfessionalReportService
from bot.pro.telegram_handlers import NOTES, PRO_STATE_KEY
from bot.pro.telegram_menu_handlers import PROJECT, TelegramProMenuHandlers


AUTHORIZED_ID = 10001
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def settings() -> Settings:
    return Settings(
        TELEGRAM_BOT_TOKEN="123456789:" + "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijk",
        OPENAI_API_KEY="sk-test-placeholder",
        SUPABASE_URL="https://example.supabase.co",
        SUPABASE_SECRET_KEY="sb_secret_test_placeholder",
        BUILDWISE_PRO_ENABLED=True,
        BUILDWISE_PRO_WEB_ENABLED=False,
        BUILDWISE_PRO_ALLOWED_TELEGRAM_IDS=str(AUTHORIZED_ID),
        BUILDWISE_PRO_DAILY_REQUEST_LIMIT=2,
        BUILDWISE_PRO_MAX_FILE_SIZE_MB=1,
    )


def structured_docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("Project Information", level=1)
    document.add_paragraph("Project: BuildWise QA Office Extension")
    document.add_paragraph("Location: Lagos, Nigeria")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Observation"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Crack near window"
    table.cell(1, 1).text = "Width not measured"
    document.save(stream)
    return stream.getvalue()


class FakeTelegramFile:
    def __init__(self, data: bytes) -> None:
        self.data = data

    async def download_as_bytearray(self):
        return bytearray(self.data)


class FakeBot:
    def __init__(self, data: bytes) -> None:
        self.data = data

    async def get_file(self, file_id: str):
        return FakeTelegramFile(self.data)


def update_with_docx(data: bytes):
    document = SimpleNamespace(
        file_name="BUILDWISE PROject.docx",
        file_id="docx-file",
        file_size=len(data),
    )
    message = SimpleNamespace(document=document, photo=[], reply_text=AsyncMock())
    update = SimpleNamespace(
        effective_message=message,
        effective_user=SimpleNamespace(id=AUTHORIZED_ID),
        effective_chat=SimpleNamespace(id=999),
    )
    return update, message


class BuildWiseProDocumentAnalysisTests(unittest.TestCase):
    def test_docx_extraction_preserves_headings_and_tables_for_analysis(self) -> None:
        data = structured_docx_bytes()
        text = ProfessionalReportService._docx_text(
            ProFile("project.docx", DOCX_MIME, data), label="primary project source"
        )
        self.assertIn("[Heading 1] Project Information", text)
        self.assertIn("[Table 1]", text)
        self.assertIn("Observation | Status", text)
        self.assertIn("Crack near window | Width not measured", text)

    def test_docx_input_explicitly_requests_analysis_not_transcription(self) -> None:
        data = structured_docx_bytes()
        content = ProfessionalReportService._file_content(
            ProFile("project.docx", DOCX_MIME, data), label="primary project source"
        )
        self.assertEqual(content["type"], "input_text")
        self.assertIn("Analyze", content["text"])
        self.assertIn("do not merely transcribe", content["text"])
        self.assertIn("gaps and uncertainties", content["text"])


class BuildWiseProDirectUploadTests(unittest.IsolatedAsyncioTestCase):
    async def test_step_one_docx_becomes_primary_source_and_moves_to_notes(self) -> None:
        data = structured_docx_bytes()
        handlers = TelegramProMenuHandlers(settings())
        context = SimpleNamespace(user_data={}, bot=FakeBot(data))
        update, message = update_with_docx(data)

        result = await handlers.project_source_direct(update, context)

        self.assertEqual(result, NOTES)
        state = context.user_data[PRO_STATE_KEY]
        self.assertEqual(state["drawing"].mime_type, DOCX_MIME)
        self.assertIn("extract only", state["project_information"])
        response = message.reply_text.await_args.args[0]
        self.assertIn("accepted for analysis", response)
        self.assertIn("tables", response)
        self.assertIn("gaps and uncertainties", response)

    async def test_invalid_docx_at_step_one_stays_on_project_step(self) -> None:
        data = b"PKnot-a-valid-docx"
        handlers = TelegramProMenuHandlers(settings())
        context = SimpleNamespace(user_data={}, bot=FakeBot(data))
        update, message = update_with_docx(data)

        result = await handlers.project_source_direct(update, context)

        self.assertEqual(result, PROJECT)
        state = context.user_data.get(PRO_STATE_KEY, {})
        self.assertNotIn("drawing", state)
        self.assertIn("valid DOCX", message.reply_text.await_args.args[0])


if __name__ == "__main__":
    unittest.main()
