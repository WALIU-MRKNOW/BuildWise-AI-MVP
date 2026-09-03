from __future__ import annotations

from io import BytesIO
import json
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock

from docx import Document

from bot.config import Settings
from bot.pro.ai_service import ProFile, ProfessionalReportService
from bot.pro.models import ProfessionalReportRequest
from bot.pro.telegram_handlers import NOTES, PRO_STATE_KEY, TEMPLATE
from bot.pro.telegram_menu_handlers import TelegramProMenuHandlers


AUTHORIZED_ID = 10001
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def settings() -> Settings:
    return Settings(
        TELEGRAM_BOT_TOKEN="123456789:" + "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijk",
        OPENAI_API_KEY="sk-test-placeholder",
        SUPABASE_URL="https://example.supabase.co",
        SUPABASE_SECRET_KEY="sb_secret_test_placeholder",
        BUILDWISE_PRO_ENABLED=True,
        BUILDWISE_PRO_WEB_ENABLED=False,
        BUILDWISE_PRO_ALLOWED_TELEGRAM_IDS=str(AUTHORIZED_ID),
        BUILDWISE_PRO_DAILY_REQUEST_LIMIT=2,
        BUILDWISE_PRO_MAX_FILE_SIZE_MB=1,
    )


def docx_bytes(text: str = "Visible cracking noted near Window W1") -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("Site Observations", level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Observation"
    table.cell(0, 1).text = "Verification"
    table.cell(1, 0).text = "Crack near window"
    table.cell(1, 1).text = "Width not measured"
    document.save(stream)
    return stream.getvalue()


class FakeTelegramFile:
    def __init__(self, data: bytes) -> None:
        self.data = data

    async def download_as_bytearray(self):
        return bytearray(self.data)


class FakeBot:
    def __init__(self, data: bytes) -> None:
        self.data = data

    async def get_file(self, file_id: str):
        return FakeTelegramFile(self.data)


def context(data: bytes):
    primary = ProFile("project.pdf", "application/pdf", b"%PDF-1.7 project")
    return SimpleNamespace(
        user_data={PRO_STATE_KEY: {"drawing": primary, "project_information": "QA Project"}},
        bot=FakeBot(data),
    )


def document_update(filename: str, data: bytes):
    document = SimpleNamespace(file_name=filename, file_id="site-file", file_size=len(data))
    message = SimpleNamespace(document=document, photo=[], text=None, reply_text=AsyncMock())
    return (
        SimpleNamespace(
            effective_message=message,
            effective_user=SimpleNamespace(id=AUTHORIZED_ID),
            effective_chat=SimpleNamespace(id=999),
        ),
        message,
    )


class BuildWiseProSiteNotesUploadTests(unittest.IsolatedAsyncioTestCase):
    async def _assert_accepted(self, filename: str, data: bytes, expected_mime: str) -> None:
        handlers = TelegramProMenuHandlers(settings())
        ctx = context(data)
        update, message = document_update(filename, data)

        result = await handlers.site_notes_file(update, ctx)

        self.assertEqual(result, TEMPLATE)
        state = ctx.user_data[PRO_STATE_KEY]
        primary = state["drawing"]
        self.assertEqual(len(primary.attachments), 1)
        self.assertEqual(primary.attachments[0].mime_type, expected_mime)
        self.assertEqual(primary.attachments[0].filename, filename)
        self.assertIn("Analyze that attached", state["site_notes"])
        first_reply = message.reply_text.await_args_list[0].args[0]
        self.assertIn("Site notes file accepted for analysis", first_reply)
        self.assertIn("will not merely read", first_reply)
        self.assertIn("gaps and uncertainties", first_reply)
        final_reply = message.reply_text.await_args_list[-1].args[0]
        self.assertIn("Step 4 of 6", final_reply)

    async def test_txt_site_notes_file_is_accepted_for_analysis(self) -> None:
        await self._assert_accepted(
            "BuildWise_Demo_Site_Notes.txt",
            b"Visible cracking near window. Crack width not measured.",
            "text/plain",
        )

    async def test_docx_site_notes_file_is_accepted_for_analysis(self) -> None:
        await self._assert_accepted("site_notes.docx", docx_bytes(), DOCX_MIME)

    async def test_pdf_site_notes_file_is_accepted_for_analysis(self) -> None:
        await self._assert_accepted(
            "site_notes.pdf",
            b"%PDF-1.7 synthetic site notes",
            "application/pdf",
        )

    async def test_image_is_not_accepted_as_site_notes_file(self) -> None:
        data = b"\x89PNG\r\n\x1a\nrest"
        handlers = TelegramProMenuHandlers(settings())
        ctx = context(data)
        update, message = document_update("site_photo.png", data)

        result = await handlers.site_notes_file(update, ctx)

        self.assertEqual(result, NOTES)
        self.assertEqual(ctx.user_data[PRO_STATE_KEY]["drawing"].attachments, ())
        rejection = message.reply_text.await_args.args[0]
        self.assertIn("Site notes must be uploaded as PDF, DOCX or TXT", rejection)
        self.assertIn("primary project source", rejection)


class FakeResponses:
    def __init__(self) -> None:
        self.kwargs = None

    async def create(self, **kwargs):
        self.kwargs = kwargs
        output = {
            "project_document_information": ["QA Project"],
            "summary": "Draft based only on supplied evidence.",
            "observations": ["Uploaded site notes state visible cracking."],
            "drawing_sketch_findings": ["Primary project source was supplied."],
            "site_note_findings": ["Crack width was not measured."],
            "relevant_technical_considerations": ["Professional verification is required."],
            "referenced_code_standard": [],
            "recommendations": ["Verify the reported condition on site."],
            "limitations_assumptions": ["No independent inspection was performed."],
            "uncertainties_missing_information": ["Crack dimensions were not supplied."],
            "review_signoff": "Draft",
            "professional_review_notice": "Draft",
            "template_structure_used": False,
            "template_sections": [],
        }
        return SimpleNamespace(output_text=json.dumps(output))


class FakeClient:
    def __init__(self) -> None:
        self.responses = FakeResponses()


class BuildWiseProSiteNotesAnalysisTests(unittest.IsolatedAsyncioTestCase):
    async def test_site_notes_attachment_is_sent_as_analysis_evidence(self) -> None:
        service = ProfessionalReportService(settings())
        fake = FakeClient()
        service._client = fake  # type: ignore[assignment]
        notes = ProFile(
            "BuildWise_Demo_Site_Notes.txt",
            "text/plain",
            b"Visible cracking near window. Width not measured.",
        )
        primary = ProFile(
            "project.pdf",
            "application/pdf",
            b"%PDF-1.7 project",
            attachments=(notes,),
        )

        await service.generate_professional_report(
            request=ProfessionalReportRequest(
                project_information="QA Project",
                site_notes="Site observations supplied as an uploaded file.",
            ),
            drawing=primary,
            template=None,
        )

        call = fake.responses.kwargs
        self.assertIsNotNone(call)
        content = call["input"][0]["content"]
        prompt = content[0]["text"]
        self.assertIn("uploaded site-notes source", prompt)
        self.assertIn("do not merely read it back", prompt)
        attachment_items = [
            item for item in content
            if item["type"] == "input_text" and "Uploaded Site Notes Source" in item["text"]
        ]
        self.assertEqual(len(attachment_items), 1)
        self.assertIn("Visible cracking near window", attachment_items[0]["text"])
        self.assertIn("do not merely transcribe", attachment_items[0]["text"])


if __name__ == "__main__":
    unittest.main()
